import io
import unittest
from unittest.mock import patch

import openpyxl
from flask import Flask, get_flashed_messages

from blueprints import ranking


class FakeCursor:
    def __init__(self):
        self.fetchone_value = None
        self.rowcount = -1
        self.institutions = []
        self.snapshot_inserts = 0
        self.lock_calls = 0
        self.executed_sql = []

    def execute(self, sql, params=None):
        normalized = " ".join(sql.split())
        self.executed_sql.append((normalized, params))
        self.fetchone_value = None
        if "pg_advisory_xact_lock" in normalized:
            self.lock_calls += 1
            self.rowcount = 1
        elif normalized.startswith("SELECT s.id"):
            self.fetchone_value = (7001,) if self.institutions else None
            self.rowcount = 1 if self.fetchone_value else 0
        elif normalized.startswith("INSERT INTO h_index_snapshots"):
            self.snapshot_inserts += 1
            self.fetchone_value = (7001,)
            self.rowcount = 1
        elif normalized.startswith("INSERT INTO h_index_institutions"):
            self.institutions.append(params)
            self.rowcount = 1
        elif normalized.startswith(("SAVEPOINT", "RELEASE SAVEPOINT")):
            pass
        else:
            raise AssertionError(f"Unexpected SQL: {normalized}")

    def fetchone(self):
        return self.fetchone_value


class FakeConnection:
    def __init__(self):
        self.commits = 0
        self.rollbacks = 0

    def commit(self):
        self.commits += 1

    def rollback(self):
        self.rollbacks += 1


def sample_workbook():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([
        "rank", "institution_name", "category", "h_scopus", "h_wos",
        "h_scholar", "national_h_index",
    ])
    for rank in range(1, 10):
        sheet.append([
            rank, f"Sample University {rank}", "university",
            100 - rank, 90 - rank, 80 - rank, 70 - rank,
        ])
    sheet.append([10, "Sample University 10", "university", "N/A", "#N/A", None, "—"])
    payload = io.BytesIO()
    workbook.save(payload)
    payload.seek(0)
    return payload


def duplicate_workbook():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([
        "rank", "institution_name", "category", "h_scopus", "h_wos",
        "h_scholar", "national_h_index",
    ])
    for rank in range(1, 10):
        sheet.append([
            rank, f"Sample University {rank}", "university",
            100 - rank, 90 - rank, 80 - rank, 70 - rank,
        ])
    sheet.append([
        10, "Sample University 9", "university", 90, 80, 70, 60,
    ])
    payload = io.BytesIO()
    workbook.save(payload)
    payload.seek(0)
    return payload


class InstitutionImportTest(unittest.TestCase):
    def setUp(self):
        self.app = Flask(__name__)
        self.app.secret_key = "test-only"
        self.cursor = FakeCursor()
        self.connection = FakeConnection()

    def upload_workbook(self, payload=None, filename="institutions-sample-10.xlsx",
                        source="staging-sample"):
        if payload is None:
            payload = sample_workbook()
        with self.app.test_request_context(
            "/admin/reyting",
            method="POST",
            data={
                "kind": "institutions",
                "snapshot_date": "2026-07-31",
                "source": source,
                "notes": "10-row isolated test",
                "file": (payload, filename),
            },
            content_type="multipart/form-data",
        ):
            ranking._handle_upload(self.cursor, self.connection)
            return get_flashed_messages(with_categories=True)

    def test_na_tokens_remain_null(self):
        for value in ("N/A", "na", "#N/A", "NULL", "None", "-", "—", None):
            with self.subTest(value=value):
                self.assertIsNone(ranking._to_int(value))

    def test_ten_row_import_is_idempotent(self):
        with patch.object(ranking, "_build_uni_list", return_value=[]), patch.object(
            ranking, "_clear_ranking_cache"
        ):
            first_messages = self.upload_workbook()
            second_messages = self.upload_workbook()

        self.assertEqual(len(self.cursor.institutions), 10)
        self.assertEqual(self.cursor.snapshot_inserts, 1)
        self.assertEqual(self.cursor.lock_calls, 2)
        self.assertEqual(self.connection.commits, 1)
        self.assertEqual(self.connection.rollbacks, 1)
        self.assertEqual(self.cursor.institutions[-1][4:8], (None, None, None, None))
        self.assertTrue(any("10 ta qator" in text for _, text in first_messages))
        self.assertTrue(any("allaqachon mavjud" in text for _, text in second_messages))
        self.assertFalse(any("h_index_scholars" in sql for sql, _ in self.cursor.executed_sql))

    def test_institution_import_is_atomic_when_a_row_is_skipped(self):
        with patch.object(ranking, "_build_uni_list", return_value=[]), patch.object(
            ranking, "_clear_ranking_cache"
        ):
            messages = self.upload_workbook(
                duplicate_workbook(),
                filename="institutions-duplicate.xlsx",
                source="staging-duplicate",
            )

        self.assertEqual(self.connection.commits, 0)
        self.assertEqual(self.connection.rollbacks, 1)
        self.assertEqual(self.cursor.snapshot_inserts, 1)
        self.assertTrue(any("qisman yozilmadi" in text for _, text in messages))


if __name__ == "__main__":
    unittest.main()
